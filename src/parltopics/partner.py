"""Excel files for the domain partner (non-coder): readable columns, clickable links, dropdowns, README sheets."""

import csv
import re
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from parltopics.config import Settings
from parltopics.sources.assemble import LangMissing, assemble
from parltopics.sources.router import AffairSource

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
EXCERPT_CHARS = 1200

PRIORITY = [
    "Priority order of the partner files:",
    "1. politmonitor_list.xlsx — now: confirm demo affairs (filter 'Keyword-free = yes').",
    "2. cv_aps_mapping.xlsx — about 20 minutes: check the pre-filled APS chapters.",
    "3. aps_gold_30.xlsx — OPTIONAL: only needed if we reach the APS step (S7).",
    "4. adjudication.xlsx — later: we fill it with disputed affairs after the evaluation run (S6).",
]


def _readme(wb: openpyxl.Workbook, title: str, lines: list[str]) -> None:
    ws = wb.active
    ws.title = "README"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    for i, line in enumerate(lines + [""] + PRIORITY, start=3):
        ws.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def _table(ws: Worksheet, headers: list[str], rows: list[list], widths: list[int], wrap: set[int] = frozenset()) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for col in wrap:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def _links(ws: Worksheet, col: int, label: str = "open") -> None:
    for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        if cell.value:
            cell.hyperlink = cell.value
            cell.value = label
            cell.style = "Hyperlink"


def _dropdown(ws: Worksheet, col: int, first_row: int, last_row: int, formula: str) -> None:
    letter = get_column_letter(col)
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.add(f"{letter}{first_row}:{letter}{last_row}")
    ws.add_data_validation(dv)
    for (cell,) in ws.iter_rows(min_row=first_row, max_row=last_row, min_col=col, max_col=col):
        cell.fill = INPUT_FILL


def excerpt(text: str) -> str:
    body = text.split("\n\n", 1)[1] if "\n\n" in text else ""
    return body[:EXCERPT_CHARS] + ("…" if len(body) > EXCERPT_CHARS else "")


def _aps(raw_dir: Path) -> tuple[list[tuple], list[tuple]]:
    wb = openpyxl.load_workbook(raw_dir / "aps_topics.xlsx", read_only=True)
    chapters = [r for r in wb["Chapters"].iter_rows(min_row=2, values_only=True) if r[0]]
    topics = [r for r in wb["Topics"].iter_rows(min_row=2, values_only=True) if r[0]]
    return chapters, topics


def politmonitor_list(settings: Settings, scan: pd.DataFrame, out: Path) -> int:
    pm = pd.read_csv(settings.data_dir / "raw" / "politmonitor_digitalisierung.csv", dtype=str, keep_default_na=False)
    flags = scan.set_index("affair_id")
    yes_no = {True: "yes", False: "no"}
    rows = []
    for r in pm.itertuples():
        f = flags.loc[int(r.openparldata_id)]
        kw = yes_no.get(f.keyword_absent, "unknown") if f.scan_error == "" else "unknown"
        tech = yes_no.get(f.tech_vocab_absent, "unknown") if f.scan_error == "" else "unknown"
        full = yes_no.get(not f.title_only, "unknown") if f.scan_error == "" else "unknown"
        rows.append([int(r.openparldata_id), r.body_key, r.number, pd.to_datetime(r.begin_date).date(), r.type_en,
                     r.title_de, r.title_fr, r.title_it, r.url, full, kw, tech, r.politmonitor_scope])
    rows.sort(key=lambda x: (x[1] != "CHE", x[1], x[3]), reverse=False)
    wb = openpyxl.Workbook()
    _readme(wb, "PolitMonitor digitalisation list (3,433 affairs)", [
        "What this is: every affair the PolitMonitor team tagged 'Digitalisierung', one row each, from the challenge "
        "file politmonitor_digitalisierung.csv (snapshot 2026-10-18). Source: OpenParlData.ch, CC BY 4.0.",
        "Use it to pick demo affairs. Use the filter arrows in the header row (sheet 'Affairs'), e.g. Parliament = CHE "
        "and Date after 2018.",
        "Column 'Keyword-free (core words)': 'yes' means none of these words appears anywhere in the affair's full "
        "text (title + texts, every language version): DE Digitalisierung, digital… ; FR numérique, numérisation, "
        "digital… ; IT digitale, digitalizzazione, digital… ('digital…' = any word starting with 'digital'). "
        "Filter it to 'yes' to see the best demo candidates: affairs a keyword search would miss.",
        "Column 'Also no tech words': 'yes' means the general tech words Daten/données/dati, elektronisch/électronique/"
        "elettronico, Cyber are absent too (stricter).",
        "Column 'Full text available': 'no' means we only have the title. Such affairs are keyword-free too easily — "
        "for demo candidates filter 'Full text available = yes' AND 'Keyword-free = yes'.",
        "'unknown' = the text could not be loaded. 'Link' opens the parliament's own page.",
        "Nothing to fill in. Please note the Affair IDs you'd like in the demo and send them back.",
    ])
    ws = wb.create_sheet("Affairs")
    _table(ws, ["Affair ID", "Parliament", "Number", "Date", "Type", "Title (DE)", "Title (FR)", "Title (IT)", "Link",
                "Full text available", "Keyword-free (core words)", "Also no tech words", "PolitMonitor scope"],
           rows, [10, 10, 16, 11, 16, 50, 50, 50, 8, 12, 14, 12, 10])
    for (cell,) in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        cell.number_format = "yyyy-mm-dd"
    _links(ws, 9)
    ws.auto_filter.ref = ws.dimensions
    wb.save(out)
    return len(rows)


def cv_aps_mapping(settings: Settings, draft_csv: Path, out: Path) -> int:
    raw_dir = settings.data_dir / "raw"
    chapters, _ = _aps(raw_dir)
    names = {c[0]: f"{c[0]} {c[3]} / {c[4]}" for c in chapters}
    with draft_csv.open(encoding="utf-8") as f:
        draft = {r["cv_topic_id"]: r for r in csv.DictReader(f)}
    vocab = pd.read_csv(raw_dir / "che_topics_vocabulary.csv", dtype=str)
    rows = []
    for v in vocab.itertuples():
        d = draft[v.topic_id]
        rows.append([v.topic_id, v.topic_de, v.topic_fr, int(v.affairs),
                     "; ".join(names[c] for c in d["aps_chapters"].split("|")), d["draft_note"], "", ""])
    wb = openpyxl.Workbook()
    _readme(wb, "Curia Vista areas → APS chapters (about 20 minutes)", [
        "What this is: the 28 Curia Vista subject areas (the federal parliament's own labels, ~70,000 affairs) and, for "
        "each, the APS chapter(s) it corresponds to. We need this to score APS predictions against Curia Vista labels.",
        "The column 'Draft APS chapter(s)' is pre-filled as a starting point. It was drafted by an AI coding assistant "
        "(Claude), not by an expert — please check every row.",
        "For each row: choose 'yes' or 'no' in 'Correct?'. If 'no', write the right chapter code(s) in 'Correction', "
        "e.g. '6.1; 4.4'. Several chapters per area are fine. If no chapter fits, write 'none'.",
        "The sheet 'APS chapters' lists all 22 chapters (DE/FR) with links. Yellow cells are for you.",
    ])
    ws = wb.create_sheet("Mapping")
    _table(ws, ["Curia Vista ID", "Curia Vista area (DE)", "Curia Vista area (FR)", "Federal affairs",
                "Draft APS chapter(s)", "Why (draft)", "Correct? (yes/no)", "Correction"],
           rows, [12, 30, 30, 10, 60, 50, 12, 30], wrap={5, 6})
    _dropdown(ws, 7, 2, len(rows) + 1, '"yes,no"')
    for (cell,) in ws.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=8, max_col=8):
        cell.fill = INPUT_FILL
    ws.auto_filter.ref = ws.dimensions
    ref = wb.create_sheet("APS chapters")
    _table(ref, ["Code", "Domain (DE)", "Domain (FR)", "Chapter (DE)", "Chapter (FR)", "Link"],
           [list(c[:6]) for c in chapters], [8, 30, 30, 32, 36, 8])
    _links(ref, 6)
    wb.save(out)
    return len(rows)


def aps_gold(settings: Settings, sample: pd.DataFrame, out: Path) -> int:
    raw_dir = settings.data_dir / "raw"
    chapters, topics = _aps(raw_dir)
    source = AffairSource(settings)
    rows = []
    for i, r in enumerate(sample.itertuples(), start=1):
        affair = source.get_affair(int(r.affair_id), r.body_key)
        try:
            src = assemble(affair, settings.max_input_chars, r.lang or None)
        except LangMissing:
            continue
        rows.append([i, int(r.affair_id), r.body_key, src.lang, src.title, src.url, excerpt(src.text),
                     "", "", "", "", "", ""])
    wb = openpyxl.Workbook()
    _readme(wb, "APS gold labels for 30 affairs (OPTIONAL)", [
        "OPTIONAL — only needed if we reach the APS step (S7). Do files 1 and 2 first.",
        "What this is: 30 affairs (15 federal, 15 cantonal; 10 each in DE, FR, IT), picked at random with a fixed seed. "
        "You label them with the APS scheme; we then compare the model to your labels. You see no model output.",
        "For each row: read the title and excerpt (or open the link for the full text). In 'APS chapter 1' pick the "
        "chapter that best describes what the affair is about, and in 'APS topic 1' the topic inside it. If a second "
        "chapter/topic is also a main subject, fill 'chapter 2' / 'topic 2'. If nothing fits, choose 'NONE' in "
        "'APS chapter 1'.",
        "Dropdowns list all 22 chapters and 165 topics (DE / FR). The sheet 'APS scheme' shows the full list.",
        "About 1–2 minutes per affair. Yellow cells are for you.",
    ])
    ws = wb.create_sheet("Affairs")
    _table(ws, ["No.", "Affair ID", "Parliament", "Language", "Title", "Link", "Text excerpt", "APS chapter 1",
                "APS topic 1", "APS chapter 2 (optional)", "APS topic 2 (optional)", "Unsure? (yes)", "Note"],
           rows, [5, 10, 10, 9, 45, 8, 80, 30, 34, 30, 34, 10, 30], wrap={5, 7})
    _links(ws, 6)
    lists = wb.create_sheet("lists")
    chapter_labels = ["NONE"] + [f"{c[0]} {c[3]} / {c[4]}" for c in chapters]
    topic_labels = [f"{t[0]} {t[5]} / {t[6]}" for t in topics]
    for j, label in enumerate(chapter_labels, start=1):
        lists.cell(row=j, column=1, value=label)
    for j, label in enumerate(topic_labels, start=1):
        lists.cell(row=j, column=2, value=label)
    lists.sheet_state = "hidden"
    last = len(rows) + 1
    for col in (8, 10):
        _dropdown(ws, col, 2, last, f"lists!$A$1:$A${len(chapter_labels)}")
    for col in (9, 11):
        _dropdown(ws, col, 2, last, f"lists!$B$1:$B${len(topic_labels)}")
    _dropdown(ws, 12, 2, last, '"yes"')
    for (cell,) in ws.iter_rows(min_row=2, max_row=last, min_col=13, max_col=13):
        cell.fill = INPUT_FILL
    ref = wb.create_sheet("APS scheme")
    _table(ref, ["Code", "Chapter (DE)", "Chapter (FR)", "Topic (DE)", "Topic (FR)", "Link"],
           [[t[0], t[3], t[4], t[5], t[6], t[7]] for t in topics], [8, 32, 36, 40, 40, 8])
    _links(ref, 6)
    wb.save(out)
    return len(rows)


def adjudication_template(settings: Settings, definition_md: Path, out: Path, rows: list[list] | None = None) -> int:
    rows = rows or []
    wb = openpyxl.Workbook()
    _readme(wb, "Review: is this affair about digitalisation? (blind)", [
        "What this is: affairs where the evaluation found a disagreement (and some where it didn't), mixed and shuffled. "
        "You don't see who said what — no model answer, no PolitMonitor tag, no hint which side flagged it. "
        "Your answer decides the 'after review' numbers.",
        "For each row: read the title and the excerpt (open the link for the full text if needed). Decide by your "
        "definition (sheet 'Definition', = data/gold/digitalisation_definition.md as committed — the rules are fixed "
        "before you start). Choose yes / no / unsure. Add a short note if useful (e.g. 'only in passing').",
        "About 1 minute per affair. Yellow cells are for you. Please don't change the other columns or the row order.",
        "This file is a template until the evaluation run (S6) fills it.",
    ])
    ws = wb.create_sheet("Review")
    _table(ws, ["Review No.", "Affair ID", "Title", "Link", "Text excerpt",
                "Is this about digitalisation? (yes/no/unsure)", "Note"],
           rows, [10, 10, 45, 8, 90, 18, 40], wrap={3, 5})
    _links(ws, 4)
    last = max(len(rows) + 1, 301)
    _dropdown(ws, 6, 2, last, '"yes,no,unsure"')
    for (cell,) in ws.iter_rows(min_row=2, max_row=last, min_col=7, max_col=7):
        cell.fill = INPUT_FILL
    definition = wb.create_sheet("Definition")
    text = re.sub(r"\*\*Author:\*\*[^.]*\.\s*", "", definition_md.read_text(encoding="utf-8"))
    lines = text.splitlines()
    for i, line in enumerate(lines, start=1):
        definition.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
    definition.column_dimensions["A"].width = 140
    wb.save(out)
    return len(rows)
