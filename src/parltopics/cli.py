import json
from pathlib import Path

import pandas as pd
import typer
import uvicorn

from parltopics.batch import run_sample, run_standard_sample
from parltopics.classify import classify_adhoc
from parltopics.config import get_settings
from parltopics.evaluate import digitalisation as digi_eval
from parltopics.evaluate import metrics_report
from parltopics.evaluate import report as reports
from parltopics.evaluate.keyword_scan import scan as keyword_scan
from parltopics.evaluate.samples import (
    build_aps_gold, build_crosslang, build_curia_vista_sample, build_demo_pool, build_digi_federal, build_first_loop, enlarge_negatives,
    federal_slice,
)
from parltopics import partner as partner_files
from parltopics.sources.router import AffairSource
from parltopics.llm import LLMClient
from parltopics.prompts import load_prompt
from parltopics.runlog import log_call, start_run
from parltopics.sources import bulk
from parltopics.sources.api import OPDClient
from parltopics.sources.assemble import LangMissing, assemble
from parltopics.store import Store
from parltopics.topics import keyword_absent, load_topic, tech_vocab_absent

RESULTS_DIR = Path("results")

app = typer.Typer(no_args_is_help=True)
data_app = typer.Typer(no_args_is_help=True, help="Challenge data files.")
app.add_typer(data_app, name="data")
sample_app = typer.Typer(no_args_is_help=True, help="Build frozen sample files.")
app.add_typer(sample_app, name="sample")
report_app = typer.Typer(no_args_is_help=True, help="Reports from stored runs.")
app.add_typer(report_app, name="report")
review_app = typer.Typer(no_args_is_help=True, help="Human review of stored predictions (append-only).")
app.add_typer(review_app, name="review")
partner_app = typer.Typer(no_args_is_help=True, help="Excel files for the domain partner.")
app.add_typer(partner_app, name="partner")
eval_app = typer.Typer(no_args_is_help=True, help="Evaluations → results/metrics.json + metrics.md.")
app.add_typer(eval_app, name="eval")
demo_app = typer.Typer(no_args_is_help=True, help="Cached demo results for the UI.")
app.add_typer(demo_app, name="demo")
check_app = typer.Typer(no_args_is_help=True, help="Live sanity checks against the model (not unit tests).")
app.add_typer(check_app, name="check")


@app.callback()
def main() -> None:
    """ParlTopics command line."""


@app.command()
def serve(host: str = "0.0.0.0", port: int = 8080) -> None:
    """Start the API server."""
    uvicorn.run("parltopics.api:app", host=host, port=port)


@app.command()
def fetch(affair_id: int, lang: str | None = typer.Option(None, help="de|fr|it, federal affairs only")) -> None:
    """Fetch one affair and print the assembled text the model will see."""
    settings = get_settings()
    affair = OPDClient(settings).get_affair(affair_id)
    try:
        source = assemble(affair, settings.max_input_chars, lang)
    except LangMissing as e:
        typer.echo(f"status=lang_missing {e}")
        raise typer.Exit(2)
    except ValueError as e:
        typer.echo(f"error: {e}")
        raise typer.Exit(2)
    typer.echo(
        f"id={affair.id} body={affair.body.key} number={affair.number} type={affair.type} "
        f"date={affair.begin_date}\nlang={source.lang} ({source.lang_source}) "
        f"detected={source.detected_lang} mismatch={source.lang_mismatch}\n"
        f"sections={','.join(source.sections)} title_only={source.title_only}\n"
        f"chars={len(source.text)}/{source.n_chars_full} truncated={source.truncated} "
        f"cut_section={source.cut_section} chars_lost={source.chars_lost}\n"
        f"url={source.url}\nsha256={source.sha256}\n---"
    )
    typer.echo(source.text)


@app.command()
def classify(
    topic: Path | None = typer.Option(None, help="topic YAML, e.g. topics/<name>.yaml (ad-hoc mode)"),
    scheme: str | None = typer.Option(None, help="standard mode: curia_vista | aps (needs --sample)"),
    affair: int | None = typer.Option(None, help="one OpenParlData affair id"),
    sample: Path | None = typer.Option(None, help="sample CSV (affair_id, body_key, lang, gold, ...)"),
    lang: str | None = typer.Option(None, help="de|fr|it, federal affairs only (with --affair)"),
    prompt_version: str | None = typer.Option(None, help="default: adhoc_v5 (ad-hoc) / standard_v1 (standard)"),
    notes: str = "",
    split: str | None = typer.Option(None, help="only rows of this split (dev|test) of the sample"),
) -> None:
    """Classify one affair (--affair) or a sample (--sample) against an ad-hoc topic (--topic) or a scheme (--scheme)."""
    settings = get_settings()
    if (affair is None) == (sample is None):
        raise typer.BadParameter("pass exactly one of --affair or --sample")
    if (topic is None) == (scheme is None):
        raise typer.BadParameter("pass exactly one of --topic (ad-hoc) or --scheme (standard)")
    store = Store(settings.db_path)
    prompt_version_adhoc = prompt_version or "adhoc_v5"
    if scheme is not None:
        if sample is None:
            raise typer.BadParameter("--scheme needs --sample")
        run = run_standard_sample(settings, store, sample, scheme, RESULTS_DIR, prompt_version or "standard_v1",
                                  notes, split)
        typer.echo(f"run_id={run.run_id} prompt_version={run.fingerprint.prompt_version[:12]}")
        return
    if sample is not None:
        run = run_sample(settings, store, sample, topic, RESULTS_DIR, prompt_version_adhoc, notes, split)
        typer.echo(f"run_id={run.run_id} prompt_version={run.fingerprint.prompt_version[:12]}")
        return
    affair_obj = OPDClient(settings).get_affair(affair)
    try:
        source = assemble(affair_obj, settings.max_input_chars, lang)
    except LangMissing as e:
        typer.echo(f"status=lang_missing {e}")
        raise typer.Exit(2)
    prompt = load_prompt(prompt_version_adhoc)
    topic_obj = load_topic(topic)
    run = start_run(RESULTS_DIR, prompt, topic_obj, topic, settings, notes=notes or f"single affair {affair}")
    store.add_run(run, topic_obj)
    store.add_source(affair_obj, source)
    prediction = classify_adhoc(
        LLMClient(settings), prompt, source, topic_obj, settings.title_only_conf_cap, run.run_id
    )
    log_call(RESULTS_DIR, prediction)
    prediction_id = store.add_prediction(prediction.model_dump() | {
        "body_key": affair_obj.body.key, "title": source.title,
        "keyword_absent": keyword_absent(source.text, topic_obj),
        "tech_vocab_absent": tech_vocab_absent(source.text, topic_obj),
    })
    store.finish_run(run.run_id, 1)
    typer.echo(prediction.model_dump_json(indent=2, exclude={"raw_json"}))
    typer.echo(f"run_id={run.run_id} prompt_version={run.fingerprint.prompt_version[:12]} prediction_id={prediction_id}")


@sample_app.command("first-loop")
def sample_first_loop(
    out: Path = Path("data/samples/first_loop.csv"),
    n_pos: int = 20,
    n_federal: int = 6,
    since: str = "2018-01-01",
    seed: int = 42,
    force: bool = False,
) -> None:
    """20 PolitMonitor positives + 20 in-scope negatives (same parliaments, since 2018), fixed seed."""
    if out.exists() and not force:
        typer.echo(f"{out} exists (frozen); use --force to rebuild")
        return
    df = build_first_loop(get_settings(), n_pos, n_federal, since, seed)
    df.to_csv(out, index=False)
    typer.echo(f"wrote {out}: {len(df)} rows, gold={df.gold.value_counts().to_dict()}")


@sample_app.command("digi-federal")
def sample_digi_federal(
    out: Path = Path("data/samples/digi_federal_2018.csv"),
    n_pos: int = 150,
    n_neg: int = 250,
    since: str = "2018-01-01",
    exclude_type: list[str] = typer.Option([], help="affair types left out of the population, e.g. Consultation"),
    pm_scope: str = typer.Option(..., help="PolitMonitor scope of the federal positives, e.g. ch"),
    dev_share: float = 0.3,
    seed: int = 42,
    exclude_sample: list[Path] = typer.Option([], help="sample files whose affairs must not be drawn (dev-only)"),
    force: bool = False,
) -> None:
    """Like-for-like digitalisation sample (docs/06 §2.3) + population counts for prevalence correction."""
    if out.exists() and not force:
        typer.echo(f"{out} exists (frozen); use --force to rebuild")
        return
    population = federal_slice(get_settings(), since, exclude_type, pm_scope)
    excluded = {int(i) for p in exclude_sample for i in pd.read_csv(p).affair_id}
    df = build_digi_federal(population, n_pos, n_neg, seed, dev_share, excluded)
    df.to_csv(out, index=False)
    meta = {"population": {"since": since, "exclude_types": exclude_type, "pm_scope": pm_scope,
                           "n": len(population), "positives": int(population.gold.sum()),
                           "negatives": int((population.gold == 0).sum()),
                           "prevalence": float(population.gold.mean())},
            "sample": {"n_pos": n_pos, "n_neg": n_neg, "seed": seed, "dev_share": dev_share,
                       "excluded_samples": [str(p) for p in exclude_sample], "excluded_affairs": len(excluded),
                       "counts": {f"{s}_{'pos' if g else 'neg'}": int(n)
                                  for (s, g), n in df.groupby(["split", "gold"]).size().items()}}}
    out.with_suffix(".meta.json").write_text(json.dumps(meta, indent=2))
    typer.echo(json.dumps(meta, indent=2))


@sample_app.command("curia-vista")
def sample_curia_vista(
    out: Path = Path("data/samples/curia_vista.csv"),
    per_area: int = 10,
    since: str = "2018-01-01",
    exclude_type: list[str] = typer.Option([], help="affair types left out, e.g. Consultation"),
    dev_share: float = 0.3,
    seed: int = 28,
    exclude_sample: list[Path] = typer.Option([], help="sample files whose affairs must not be drawn"),
    force: bool = False,
) -> None:
    """Stratified Curia Vista sample (docs/06 §3.1): >= per_area affairs per subject area, dev/test grouped by title."""
    if out.exists() and not force:
        typer.echo(f"{out} exists (frozen); use --force to rebuild")
        return
    population = federal_slice(get_settings(), since, exclude_type, "ch")
    excluded = {int(i) for p in exclude_sample for i in pd.read_csv(p).affair_id}
    df = build_curia_vista_sample(population, per_area, seed, dev_share, excluded)
    df.to_csv(out, index=False)
    typer.echo(f"wrote {out}: {len(df)} affairs, splits {df.split.value_counts().to_dict()}, "
               f"mean gold labels {df.gold.str.split('|').map(len).mean():.2f}")


@sample_app.command("crosslang")
def sample_crosslang(
    out: Path = Path("data/samples/crosslang.csv"),
    n_pos: int = 50,
    n_neg: int = 50,
    since: str = "2018-01-01",
    exclude_type: list[str] = typer.Option([], help="affair types left out, e.g. Consultation"),
    pm_scope: str = typer.Option(..., help="PolitMonitor scope of federal positives, e.g. ch"),
    seed: int = 3,
    force: bool = False,
) -> None:
    """Cross-language sample (docs/06 §4): federal affairs complete in every language, one row per language."""
    if out.exists() and not force:
        typer.echo(f"{out} exists (frozen); use --force to rebuild")
        return
    settings = get_settings()
    df = build_crosslang(settings, federal_slice(settings, since, exclude_type, pm_scope), n_pos, n_neg, seed, set())
    df.to_csv(out, index=False)
    typer.echo(f"wrote {out}: {df.affair_id.nunique()} affairs × {df.lang.nunique()} languages = {len(df)} rows; "
               f"gold {df.drop_duplicates('affair_id').gold.value_counts().to_dict()}")


@sample_app.command("enlarge-negatives")
def sample_enlarge(
    base: Path = Path("data/samples/digi_federal_2018.csv"),
    out: Path = Path("data/samples/digi_federal_2018_v2.csv"),
    n_new: int = 600,
    dev_share: float = 0.58,
    seed: int = 4242,
    exclude_sample: list[Path] = typer.Option([], help="sample files whose affairs must not be drawn"),
    force: bool = False,
) -> None:
    """Add untagged in-scope affairs to a frozen sample (D42): existing rows keep their split."""
    if out.exists() and not force:
        typer.echo(f"{out} exists (frozen); use --force to rebuild")
        return
    meta = json.loads(base.with_suffix(".meta.json").read_text())
    pop = meta["population"]
    population = federal_slice(get_settings(), pop["since"], pop["exclude_types"], pop["pm_scope"])
    existing = pd.read_csv(base, dtype={"lang": str}, keep_default_na=False)
    excluded = {int(i) for p in exclude_sample for i in pd.read_csv(p).affair_id}
    df = enlarge_negatives(existing, population, n_new, seed, dev_share, excluded)
    df.to_csv(out, index=False)
    meta["sample"] |= {"enlarged_from": str(base), "n_new_negatives": n_new, "enlarge_seed": seed,
                       "enlarge_dev_share": dev_share,
                       "counts": {f"{s}_{'pos' if g else 'neg'}": int(n)
                                  for (s, g), n in df.groupby(["split", "gold"]).size().items()}}
    out.with_suffix(".meta.json").write_text(json.dumps(meta, indent=2))
    typer.echo(json.dumps(meta["sample"], indent=2))


@sample_app.command("demo-pool")
def sample_demo_pool(
    out: Path = Path("data/samples/demo_pool.csv"),
    since: str = "2018-01-01",
    seed: int = 7,
    n_keyword_free: int = 60,
    n_other_positive: int = 40,
    n_negative: int = 100,
    force: bool = False,
) -> None:
    """Fixed demo pool for the UI (docs/02 §4); needs the PolitMonitor keyword scan (parltopics partner files)."""
    if out.exists() and not force:
        typer.echo(f"{out} exists (frozen); use --force to rebuild")
        return
    settings = get_settings()
    scan = pd.read_csv(settings.data_dir / "cache" / "politmonitor_keyword_scan.csv", keep_default_na=False)
    df = build_demo_pool(settings, scan, since, seed, n_keyword_free, n_other_positive, n_negative)
    df.to_csv(out, index=False)
    typer.echo(f"wrote {out}: {len(df)} rows; {df.split.value_counts().to_dict()}; "
               f"federal {int((df.body_key == 'CHE').sum())}, parliaments {df.body_key.nunique()}")


@report_app.command("first-loop")
def report_first_loop(run_id: str, compare_to: str | None = None) -> None:
    """Table, summary, truncation report; optionally what changed vs another run on the same sample."""
    store = Store(get_settings().db_path)
    manifest, rows = reports.load_run(store, run_id)
    typer.echo(reports.table(rows))
    typer.echo("")
    typer.echo(reports.summary_text(manifest, rows))
    typer.echo("")
    typer.echo(reports.truncation_report(rows))
    if compare_to:
        old_manifest, old_rows = reports.load_run(store, compare_to)
        old_fp, new_fp = old_manifest["fingerprint"]["parts"], manifest["fingerprint"]["parts"]
        changed = [k for k in new_fp if old_fp.get(k) != new_fp[k]]
        typer.echo(f"\nvs run {compare_to} (prompt parts changed: {changed})")
        typer.echo(reports.compare(old_rows, rows))


@app.command()
def status() -> None:
    """Where the project stands: build slices, success criteria, submission checklist (from real artefacts)."""
    from parltopics.status import collect, render

    typer.echo(render(collect(Path.cwd())))


@app.command()
def runs(limit: int = 20) -> None:
    """List stored runs, newest first."""
    for r in Store(get_settings().db_path).runs()[:limit]:
        typer.echo(f"{r['id']}  {r['model']:<18} {r['prompt_template']} {r['prompt_version'][:12]}  "
                   f"{r['topic_name']:<15} n={r['n_predictions']:<4} finished={bool(r['finished_at'])}  {r['notes']}")


@review_app.command("add")
def review_add(
    prediction_id: int,
    decision: str = typer.Argument(..., help="confirm | reject | correct"),
    label: str | None = typer.Option(None, help="corrected label (required for correct)"),
    reviewer: str = "",
    note: str = "",
) -> None:
    """Record a human review. The prediction itself is never changed."""
    review_id = Store(get_settings().db_path).add_review(prediction_id, decision, label, reviewer, note)
    typer.echo(f"review_id={review_id}")


@review_app.command("show")
def review_show(prediction_id: int) -> None:
    """Show a prediction's reviews (oldest first; the latest one counts)."""
    for r in Store(get_settings().db_path).reviews(prediction_id):
        typer.echo(f"{r['id']} {r['created_at']} {r['decision']} {r['corrected_label'] or ''} "
                   f"by={r['reviewer']} {r['note']}")


@eval_app.command("digitalisation")
def eval_digitalisation(
    run_ids: list[str] = typer.Argument(..., help="one run, or a dev run + a test run with the same prompt_version"),
    sample: Path = Path("data/samples/digi_federal_2018.csv"),
    topic: Path = Path("topics/digitalisation.yaml"),
    cv_area: list[str] = typer.Option(..., help="Curia Vista area id(s) of the baseline, e.g. 34"),
    thresholds: str = "0.5,0.6,0.7,0.8,0.9",
    n_boot: int = 1000,
    seed: int = 42,
    review_budget: int = 40,
) -> None:
    """Like-for-like eval (docs/06 §2): baselines, test metrics, dev threshold, blind review sheet."""
    settings = get_settings()
    store = Store(settings.db_path)
    grid = [float(t) for t in thresholds.split(",")]
    m = digi_eval.evaluate(settings, store, run_ids, sample, load_topic(topic), cv_area, grid, n_boot, seed)
    metrics_report.update_metrics_json(RESULTS_DIR / "metrics.json", "digitalisation_federal_2018", m)
    md = metrics_report.digitalisation_markdown(m)
    write_metrics_md()
    rows, key = digi_eval.review_rows(store, run_ids, sample, m["threshold"]["chosen_on_dev"], review_budget, seed,
                                      partner_files.excerpt)
    if settings.review_lang:
        rows = partner_files.in_review_lang(settings, rows, key, settings.review_lang)
    partner_files.adjudication_template(settings, Path("data/gold/digitalisation_definition.md"),
                                        Path("data/gold/for_partner/adjudication.xlsx"), rows)
    pd.DataFrame(key).to_csv(RESULTS_DIR / "adjudication_key.csv", index=False)
    typer.echo(md)
    typer.echo(f"review sheet: {len(rows)} affairs → data/gold/for_partner/adjudication.xlsx "
               f"(key: results/adjudication_key.csv)")


@check_app.command("topic")
def check_topic(
    name: str,
    definition: str = "",
    max_matches: int = typer.Option(..., help="fail if more affairs than this are shown as matches"),
    pool: Path = Path("data/samples/demo_pool.csv"),
) -> None:
    """Classify the demo pool with a custom topic exactly as the UI does; list shown matches; fail above the limit."""
    from parltopics.api import SearchIn, TopicIn, build_topic, match_threshold
    from parltopics.batch import iter_classify

    settings = get_settings()
    store = Store(settings.db_path)
    topic, ref = build_topic(SearchIn(topic=TopicIn(name=name, definition=definition)))
    prompt = load_prompt("adhoc_v5")
    run = start_run(RESULTS_DIR, prompt, topic, Path(ref), settings, sample_file=str(pool), notes=f"check topic {name}")
    store.add_run(run, topic)
    rows = pd.read_csv(pool, dtype={"lang": str}, keep_default_na=False).to_dict("records")
    records = list(iter_classify(settings, store, RESULTS_DIR, run, topic, prompt, rows))
    store.finish_run(run.run_id, len(records))
    threshold = match_threshold(settings)
    shown = [r for r in records if r["status"] == "ok" and r["label"] == "match" and r["confidence"] >= threshold]
    labels = pd.Series([r.get("label") or r["status"] for r in records]).value_counts().to_dict()
    typer.echo(f"{name!r} (definition: {topic.definition!r}) run {run.run_id}: {len(records)} affairs, labels {labels}, "
               f"shown at t={threshold}: {len(shown)}")
    for r in shown:
        typer.echo(f"  {r['affair_id']} {r['body_key']} {r['confidence']:.2f} | {r['title'][:70]} | {r['evidence'][:90]}")
    if len(shown) > max_matches:
        typer.echo(f"FAIL: {len(shown)} > {max_matches}")
        raise typer.Exit(1)
    typer.echo("PASS")


@demo_app.command("export")
def demo_export(pool: Path = Path("data/samples/demo_pool.csv")) -> None:
    """Export the latest finished demo-pool run of every preset topic to results/demo/<preset>.json (committed)."""
    from parltopics import api as web
    from parltopics import demo

    store = Store(get_settings().db_path)
    n_pool = len(pd.read_csv(pool))
    for preset_id, (_, topic) in web.presets().items():
        run = store.latest_run(web.topic_sha(topic), str(pool), n_pool)
        if run is None:
            typer.echo(f"{preset_id}: no finished run on {pool}")
            continue
        n = demo.export_run(store, run["id"], RESULTS_DIR / "demo" / f"{preset_id}.json")
        typer.echo(f"{preset_id}: run {run['id']} ({run['prompt_template']}, {n} predictions) → results/demo/{preset_id}.json")
    for scheme_id in web.SCHEMES:
        run = store.latest_run(web.scheme_sha(scheme_id), str(pool), n_pool)
        if run is None:
            typer.echo(f"scheme {scheme_id}: no finished run on {pool}")
            continue
        n = demo.export_run(store, run["id"], RESULTS_DIR / "demo" / f"scheme_{scheme_id}.json")
        typer.echo(f"scheme {scheme_id}: run {run['id']} ({n} label records) → results/demo/scheme_{scheme_id}.json")


@eval_app.command("curia-vista")
def eval_curia_vista(
    run_ids: list[str] = typer.Argument(..., help="runs with the same prompt_version (e.g. dev, or dev + test)"),
    sample: Path = Path("data/samples/curia_vista.csv"),
    n_boot: int = 1000,
    seed: int = 28,
) -> None:
    """Curia Vista standard mode (docs/06 §3.1): at-least-one-correct, micro/macro F1, per area, cardinality."""
    from parltopics.evaluate.multilabel import multilabel_summary
    from parltopics.schemes import load_curia_vista

    settings = get_settings()
    store = Store(settings.db_path)
    versions = {store.run(r)["prompt_version"] for r in run_ids}
    if len(versions) != 1:
        raise typer.BadParameter(f"runs have different prompt_versions: {versions}")
    smp = pd.read_csv(sample, dtype={"lang": str, "gold": str}, keep_default_na=False).set_index("affair_id")
    rows = [r for run_id in run_ids for r in store.predictions(run_id)]
    excluded = {a for a in {r["affair_id"] for r in rows}
                if any(r["affair_id"] == a and (r["status"] == "lang_missing" or r.get("lang_mismatch")) for r in rows)}
    labels = [c.source_id for c in load_curia_vista(settings.data_dir / "raw").categories]
    out = {"run_ids": run_ids, "prompt_version": versions.pop(), "model": store.run(run_ids[0])["model"],
           "excluded_lang_missing_or_mismatch": len(excluded), "status_counts": {}}
    for r in rows:
        out["status_counts"][r["status"]] = out["status_counts"].get(r["status"], 0) + 1
    for split in ("dev", "test"):
        ids = [a for a in smp.index[smp.split == split] if a not in excluded and a in {r["affair_id"] for r in rows}]
        if not ids:
            continue
        gold = {a: set(smp.loc[a, "gold"].split("|")) for a in ids}
        pred: dict[int, set[str]] = {a: set() for a in ids}
        for r in rows:
            if r["affair_id"] in pred and r["status"] == "ok" and r["label"]:
                pred[r["affair_id"]].add(r["label"])
        out[split] = multilabel_summary(gold, pred, labels, n_boot, seed)
        s = out[split]
        typer.echo(f"[{split}] n={s['n']}  at-least-one-correct {s['at_least_one_correct']['correct']}/{s['n']} = "
                   f"{s['at_least_one_correct']['rate']:.1%} (CI {s['at_least_one_correct']['ci95'][0]:.1%}–"
                   f"{s['at_least_one_correct']['ci95'][1]:.1%})  micro-F1 {s['micro']['f1']:.3f} "
                   f"(CI {s['micro']['f1_ci95'][0]:.3f}–{s['micro']['f1_ci95'][1]:.3f}; P {s['micro']['precision']:.3f} "
                   f"R {s['micro']['recall']:.3f})  macro-F1 {s['macro_f1']:.3f}  labels/affair gold "
                   f"{s['cardinality']['gold_mean']:.2f} pred {s['cardinality']['pred_mean']:.2f}  "
                   f"NONE_OTHER {s['none_other_rate']:.1%}")
    typer.echo(f"statuses {out['status_counts']}; excluded (lang missing/mismatch) {len(excluded)}")
    metrics_report.update_metrics_json(RESULTS_DIR / "metrics.json", "curia_vista_standard", out)
    write_metrics_md()


def write_metrics_md() -> None:
    metrics = json.loads((RESULTS_DIR / "metrics.json").read_text())
    (RESULTS_DIR / "metrics.md").write_text(metrics_report.full_markdown(metrics))


@eval_app.command("crosslang")
def eval_crosslang(
    adhoc_run: str,
    standard_run: str,
    sample: Path = Path("data/samples/crosslang.csv"),
) -> None:
    """Cross-language consistency (docs/06 §4) from one ad-hoc and one standard run on the crosslang sample."""
    from parltopics.api import match_threshold
    from parltopics.evaluate.crosslang import crosslang_binary, crosslang_sets

    settings = get_settings()
    store = Store(settings.db_path)
    smp = pd.read_csv(sample, dtype={"lang": str}, keep_default_na=False)
    langs = sorted(smp.lang.unique())
    gold = smp.drop_duplicates("affair_id").set_index("affair_id").gold.to_dict()
    threshold = match_threshold(settings)
    usable = lambda r: r["status"] not in ("lang_missing", "fetch_error", "parse_error") and not r.get("lang_mismatch")
    binary: dict[int, dict[str, int]] = {}
    for r in store.predictions(adhoc_run):
        if usable(r):
            shown = r["status"] == "ok" and r["label"] == "match" and (r["confidence"] or 0) >= threshold
            binary.setdefault(r["affair_id"], {})[r["lang"]] = int(shown)
    sets: dict[int, dict[str, list]] = {}
    for r in sorted(store.predictions(standard_run), key=lambda r: (r["affair_id"], r["lang"], r.get("rank") or 9)):
        if usable(r):
            labs = sets.setdefault(r["affair_id"], {}).setdefault(r["lang"], [])
            if r["status"] == "ok" and r["label"] and r["label"] != "NONE_OTHER":
                labs.append(r["label"])
    adhoc = crosslang_binary(binary, langs)
    recall = {lang: (sum(v.get(lang, 0) for a, v in binary.items() if gold.get(a) == 1)
                     / max(1, sum(1 for a, v in binary.items() if gold.get(a) == 1 and lang in v))) for lang in langs}
    out = {"sample": str(sample), "threshold": threshold, "runs": [adhoc_run, standard_run], "adhoc": adhoc,
           "standard": crosslang_sets(sets, langs), "recall_by_lang": recall}
    metrics_report.update_metrics_json(RESULTS_DIR / "metrics.json", "cross_language", out)
    write_metrics_md()
    typer.echo(metrics_report.crosslang_markdown(out))


@eval_app.command("adjudication")
def eval_adjudication(
    answers: Path = Path("data/gold/adjudication.xlsx"),
    key: Path = Path("results/adjudication_key.csv"),
    exclude_affair: list[int] = typer.Option([65590], help="affairs seen outside the blind setting (D58)"),
    n_boot: int = 2000,
    seed: int = 42,
) -> None:
    """After-review metrics from the reviewer's sheet (D55 rule, D56 readings, D58 exclusions)."""
    import openpyxl

    from parltopics.evaluate.adjudication import after_review

    ws = openpyxl.load_workbook(answers, data_only=True)["Review"]
    head = [c.value for c in ws[1]]
    sheet = pd.DataFrame([r[:len(head)] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[1]], columns=head)
    rows = pd.read_csv(key).merge(sheet, left_on="review_no", right_on="Review No.")
    if (rows["affair_id"] != rows["Affair ID"]).any():
        raise typer.BadParameter("sheet rows do not match the key (different sheet version)")
    norm = lambda v: "" if pd.isna(v) else str(v).strip().lower()
    rows["digital"], rows["goal"] = rows["Digitale Technik drin?"].map(norm), rows["Ziel oder Nebensache?"].map(norm)
    rows["note"] = rows["Note"].map(norm)
    rows["cell"] = rows.apply(lambda r: ("TP" if r.politmonitor == 1 else "TN") if r.group == "agree"
                              else ("FP" if r.group == "model_yes_pm_no" else "FN"), axis=1)
    answered = rows[rows.digital != ""]
    excluded = answered[answered.affair_id.isin(exclude_affair) | answered.note.str.contains("bekannt")]
    usable = answered.drop(excluded.index)
    readings = {
        "strict": lambda r: 1 if r.digital == "ja" and r.goal == "ziel" else 0 if r.digital == "nein" or r.goal == "nebensache" else None,
        "broad": lambda r: 1 if r.digital == "ja" else 0 if r.digital == "nein" else None,
    }
    metrics = json.loads((RESULTS_DIR / "metrics.json").read_text())
    digi = metrics["digitalisation_federal_2018"]
    t = digi["test"]["at_chosen_threshold"]
    sizes = {"TP": t["tp"], "FN": t["fn"], "FP": t["fp"], "TN": t["tn"]}
    out = {"answers_file": str(answers), "rule": "D55", "rows": len(rows), "answered": len(answered),
           "excluded": {"affair_ids": excluded.affair_id.tolist(), "reason": "seen outside blind setting / marked bekannt"},
           "before": {"precision": t["precision_corrected"], "precision_ci95": t["precision_corrected_ci95"],
                      "recall": t["recall"], "recall_ci95": t["recall_ci95"]}, "readings": {}}
    for name, fn in readings.items():
        lab = usable.assign(y=usable.apply(fn, axis=1))
        decided = lab[lab.y.notna()]
        ans = {c: decided[decided.cell == c].y.astype(int).tolist() for c in ("TP", "FN", "FP", "TN")}
        if not all(ans.values()):
            typer.echo(f"{name}: a cell has no decided answers — cannot estimate")
            continue
        res = after_review(sizes, ans, digi["population"]["prevalence"], n_boot, seed)
        res["unsure"] = int(lab.y.isna().sum())
        res["model_found_pm_missed_confirmed"] = decided[(decided.cell == "FP") & (decided.y == 1)][["affair_id", "Title"]].to_dict("records")
        res["real_misses"] = decided[(decided.cell == "FN") & (decided.y == 1)][["affair_id", "Title"]].to_dict("records")
        out["readings"][name] = res
        typer.echo(f"[{name}] precision {res['precision']:.1%} ({res['precision_ci95'][0]:.1%}–{res['precision_ci95'][1]:.1%})  "
                   f"recall {res['recall']:.1%} ({res['recall_ci95'][0]:.1%}–{res['recall_ci95'][1]:.1%})  "
                   f"answered per cell {res['answered']}  reviewer-yes rate "
                   + ", ".join(f"{c} {v:.0%}" for c, v in res["reviewer_yes_rate"].items()) + f"  unsure {res['unsure']}")
    typer.echo(f"before review: precision {t['precision_corrected']:.1%} recall {t['recall']:.1%}; "
               f"answered {len(answered)}/{len(rows)}, excluded {excluded.affair_id.tolist()}")
    metrics_report.update_metrics_json(RESULTS_DIR / "metrics.json", "adjudication", out)
    write_metrics_md()


@eval_app.command("cost")
def eval_cost(
    run_ids: list[str],
    concurrency_used: int = typer.Option(..., help="LLM_CONCURRENCY used for runs that did not record it"),
    target_affairs: int = 300_000,
) -> None:
    """Cost/time per affair and the scale-up estimate (docs/06 §6) from stored runs."""
    from parltopics.evaluate.cost import cost_time

    out = cost_time(Store(get_settings().db_path), run_ids, concurrency_used, target_affairs, [4, 8, 16, 32])
    metrics_report.update_metrics_json(RESULTS_DIR / "metrics.json", "cost_time", out)
    write_metrics_md()
    typer.echo(metrics_report.cost_markdown(out))


@partner_app.command("files")
def partner_build(
    out_dir: Path = Path("data/gold/for_partner"),
    topic: Path = Path("topics/digitalisation.yaml"),
    aps_sample: Path = Path("data/samples/aps_gold.csv"),
) -> None:
    """politmonitor_list, cv_aps_mapping, aps_gold_30 and the adjudication template (all with README sheets)."""
    settings = get_settings()
    out_dir.mkdir(parents=True, exist_ok=True)
    scan_cache = settings.data_dir / "cache" / "politmonitor_keyword_scan.csv"
    if not scan_cache.exists():
        pm = pd.read_csv(settings.data_dir / "raw" / "politmonitor_digitalisierung.csv", dtype=str)
        keyword_scan(AffairSource(settings), pm, load_topic(topic), settings.llm_concurrency).to_csv(scan_cache, index=False)
    scan = pd.read_csv(scan_cache, keep_default_na=False)
    for col in ("keyword_absent", "tech_vocab_absent", "title_only"):
        scan[col] = scan[col].astype(str).map({"True": True, "False": False})
    n = partner_files.politmonitor_list(settings, scan, out_dir / "politmonitor_list.xlsx")
    typer.echo(f"politmonitor_list.xlsx: {n} rows, keyword-free {int(scan.keyword_absent.eq(True).sum())}, "
               f"unknown {int((scan.scan_error != '').sum())}")
    n = partner_files.cv_aps_mapping(settings, Path("data/gold/cv_aps_mapping_draft.csv"), out_dir / "cv_aps_mapping.xlsx")
    typer.echo(f"cv_aps_mapping.xlsx: {n} rows")
    if not aps_sample.exists():
        first_loop = set(pd.read_csv("data/samples/first_loop.csv").affair_id)
        build_aps_gold(settings, 5, 5, "2018-01-01", 7, first_loop).to_csv(aps_sample, index=False)
    n = partner_files.aps_gold(settings, pd.read_csv(aps_sample, dtype={"lang": str}, keep_default_na=False),
                               out_dir / "aps_gold_30.xlsx")
    typer.echo(f"aps_gold_30.xlsx: {n} rows")
    partner_files.adjudication_template(settings, Path("data/gold/digitalisation_definition.md"),
                                        out_dir / "adjudication.xlsx")
    typer.echo("adjudication.xlsx: template")


@data_app.command("load-bulk")
def data_load_bulk() -> None:
    """Load federal affairs + texts from data/raw into data/cache/bulk.db (once; > ~500 affairs, D11)."""
    settings = get_settings()
    counts = bulk.load_bulk(settings.data_dir / "raw", settings.bulk_db_path)
    typer.echo(f"loaded {counts} into {settings.bulk_db_path}")


@data_app.command("download")
def data_download(force: bool = False) -> None:
    """Download challenge files into data/raw and compare with committed checksums."""
    settings = get_settings()
    raw_dir = settings.data_dir / "raw"
    checksum_file = settings.data_dir / "raw_files.sha256"
    for name in bulk.download_challenge_files(raw_dir, force):
        typer.echo(f"downloaded {name}")
    expected = bulk.read_checksums(checksum_file)
    status = bulk.check_files(raw_dir, expected)
    for name, state in status.items():
        typer.echo(f"{state:8} {name}")
    if not expected:
        checksum_file.write_text(
            "".join(f"{bulk.sha256_file(raw_dir / n)}  {n}\n" for n in bulk.CHALLENGE_FILES)
        )
        typer.echo(f"wrote {checksum_file}")
    if any(s in {"missing", "changed"} for s in status.values()):
        raise typer.Exit(1)
